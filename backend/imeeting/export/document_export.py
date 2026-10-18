"""Word / PDF 导出（含公司页眉、页脚）。"""

from __future__ import annotations

import io
from datetime import datetime
from xml.sax.saxutils import escape


def _mmss(ms: int) -> str:
    total = max(0, ms) // 1000
    return f"{total // 60:02d}:{total % 60:02d}"


def _minutes_dict(minutes) -> dict:
    if minutes is None:
        return {"summary": "", "topics": [], "todos": []}
    import json

    def load(raw):
        try:
            value = json.loads(raw or "[]")
            return value if isinstance(value, list) else []
        except Exception:
            return []

    return {
        "summary": minutes.summary,
        "topics": load(minutes.topics),
        "todos": load(minutes.todos),
    }


def build_docx(meeting, segments, minutes, *, company="", footer="") -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(10.5)

    title = doc.add_heading(meeting.title, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta = []
    if meeting.scheduled_at:
        meta.append(
            f"会议时间：{meeting.scheduled_at.strftime('%Y-%m-%d %H:%M')}"
        )
    if meeting.participants:
        meta.append(f"参会人：{meeting.participants}")
    if meeting.meeting_type:
        meta.append(f"类型：{meeting.meeting_type}")
    if meeting.agenda:
        meta.append(f"议程：{meeting.agenda}")
    for line in meta:
        p = doc.add_paragraph(line)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    data = _minutes_dict(minutes)
    if data["summary"]:
        doc.add_heading("会议摘要", level=1)
        doc.add_paragraph(data["summary"])
    if data["topics"]:
        doc.add_heading("议题与结论", level=1)
        for topic in data["topics"]:
            doc.add_heading(str(topic.get("title") or "议题"), level=2)
            if topic.get("detail"):
                doc.add_paragraph(str(topic["detail"]))
            if topic.get("conclusion"):
                doc.add_paragraph(f"结论：{topic['conclusion']}")
    if data["todos"]:
        doc.add_heading("待办事项", level=1)
        for todo in data["todos"]:
            owner = f"（{todo.get('owner')}）" if todo.get("owner") else ""
            due = f" 截止：{todo.get('due')}" if todo.get("due") else ""
            doc.add_paragraph(
                f"• {todo.get('task', '')}{owner}{due}", style="List Bullet"
            )

    doc.add_heading("完整转写", level=1)
    for seg in segments:
        speaker = f"{seg.speaker}：" if seg.speaker else ""
        doc.add_paragraph(f"[{_mmss(seg.start_ms)}] {speaker}{seg.text}")

    section = doc.sections[0]
    if company:
        header = section.header.paragraphs[0]
        header.text = company
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if footer:
        section.footer.paragraphs[0].text = footer

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def build_pdf(meeting, segments, minutes, *, company="", footer="") -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
    )

    font_path = r"C:\Windows\Fonts\msyh.ttc"
    try:
        pdfmetrics.registerFont(TTFont("MSYH", font_path, subfontIndex=0))
        font_name = "MSYH"
    except Exception:
        font_name = "Helvetica"

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "Chtitle",
        parent=styles["Title"],
        fontName=font_name,
        fontSize=18,
        leading=24,
        spaceAfter=12,
    )
    h1 = ParagraphStyle(
        "Chh1", parent=styles["Heading1"], fontName=font_name, fontSize=14, leading=20
    )
    h2 = ParagraphStyle(
        "Chh2", parent=styles["Heading2"], fontName=font_name, fontSize=12, leading=18
    )
    body = ParagraphStyle(
        "Chbody", parent=styles["BodyText"], fontName=font_name, fontSize=10.5, leading=16
    )
    small = ParagraphStyle(
        "Chsmall", parent=body, fontSize=9, leading=13, textColor=colors.grey
    )

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=42,
        rightMargin=42,
        topMargin=48,
        bottomMargin=42,
    )
    story = [Paragraph(escape(meeting.title), title_style)]
    meta = []
    if meeting.scheduled_at:
        meta.append(f"会议时间：{meeting.scheduled_at.strftime('%Y-%m-%d %H:%M')}")
    if meeting.participants:
        meta.append(f"参会人：{meeting.participants}")
    if meeting.agenda:
        meta.append(f"议程：{meeting.agenda}")
    if meta:
        story.append(Paragraph("<br/>".join(escape(x) for x in meta), small))
        story.append(Spacer(1, 10))

    data = _minutes_dict(minutes)
    if data["summary"]:
        story.append(Paragraph("会议摘要", h1))
        story.append(Paragraph(escape(data["summary"]), body))
    if data["topics"]:
        story.append(Paragraph("议题与结论", h1))
        for topic in data["topics"]:
            story.append(Paragraph(escape(str(topic.get("title") or "议题")), h2))
            if topic.get("detail"):
                story.append(Paragraph(escape(str(topic["detail"])), body))
            if topic.get("conclusion"):
                story.append(
                    Paragraph(f"结论：{escape(str(topic['conclusion']))}", body)
                )
    if data["todos"]:
        story.append(Paragraph("待办事项", h1))
        for todo in data["todos"]:
            owner = f"（{todo.get('owner')}）" if todo.get("owner") else ""
            due = f" 截止：{todo.get('due')}" if todo.get("due") else ""
            story.append(
                Paragraph(
                    f"• {escape(str(todo.get('task', '')))}{escape(owner)}"
                    f"{escape(due)}",
                    body,
                )
            )

    story.append(PageBreak())
    story.append(Paragraph("完整转写", h1))
    for seg in segments:
        speaker = f"{seg.speaker}：" if seg.speaker else ""
        story.append(
            Paragraph(
                f"[{_mmss(seg.start_ms)}] {escape(speaker)}"
                f"{escape(seg.text or '')}",
                body,
            )
        )

    def _footer(canvas, _doc):
        canvas.saveState()
        canvas.setFont(font_name, 8)
        canvas.setFillColor(colors.grey)
        text = footer or company or ""
        if text:
            canvas.drawString(42, 24, text[:120])
        canvas.drawRightString(A4[0] - 42, 24, str(datetime.now().year))
        canvas.restoreState()

    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return buffer.getvalue()
